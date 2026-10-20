# -*- coding: utf-8 -*-
"""PLANTILLAS DE LIBROS DE COMPRAS Y VENTAS  ->  assets/plantillas/

    python herramientas/plantillas_libros.py 2025 2026

Una hoja de instrucciones y doce hojas de mes, todas con el mismo formato.
Los encabezados son EXACTAMENTE los que reconoce el importador de la app
(assets/importador.js, COLUMNAS): si se cambia un encabezado aquí, el
importador deja de encontrar esa columna. El período de cada hoja va en la
fila 3 («Mes Enero: 01/01/2026 al 31/01/2026»), que es de donde el
importador lo lee; para una quincena basta cambiar ese rango.
"""
import calendar
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, CYAN, SLATE, CLARO, AMARILLO = '003057', '00BFFF', '545E68', 'EEF4F8', 'FFF7D6'
F = 'Arial'
MESES = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO',
         'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']
FILAS = 300          # renglones por mes
DESDE = 6            # primera fila de datos

# (encabezado, ancho, formato, qué va, ejemplo)
VENTAS = [
    ('Nº de Operación', 9, '0', 'Correlativo de la fila', '1'),
    ('Fecha de la Factura', 12, 'DD/MM/YYYY', 'Fecha de emisión', '05/01/2026'),
    ('Nº R.I.F.', 14, '@', 'RIF del cliente', 'J-12345678-9'),
    ('Nombre o Razón Social', 34, '@', 'Nombre del cliente', 'CLIENTE EJEMPLO, C.A.'),
    ('N° de Factura', 12, '@', 'Número de la factura', '000101'),
    ('N° Control de Factura', 14, '@', 'Número de control', '00-000101'),
    ('N° Nota de Débito', 12, '@', 'Solo si la fila es una nota de débito', ''),
    ('N° Nota de Crédito', 12, '@', 'Solo si la fila es una nota de crédito', ''),
    ('N° Factura Afectada', 12, '@', 'Factura a la que se aplica la nota', ''),
    ('Total Ventas Incluyendo el IVA', 16, '#,##0.00', 'Total del documento', '1.160,00'),
    ('Ventas Internas No Gravadas', 15, '#,##0.00', 'Monto exento o no gravado', '0,00'),
    ('Base Imponible', 15, '#,##0.00', 'Base gravada', '1.000,00'),
    ('% Alícuota', 9, '0%', 'Alícuota del IVA (16%)', '16%'),
    ('Impuesto IVA', 14, '#,##0.00', 'IVA de la factura', '160,00'),
    ('IGTF', 12, '#,##0.00', 'IGTF cobrado (pagos en divisas)', '0,00'),
    ('Comprobante de Retención', 17, '@', 'N° del comprobante que te entregó el cliente', ''),
    ('IVA Retenido por el Comprador', 15, '#,##0.00', 'IVA que te retuvo el cliente', '0,00'),
]
COMPRAS = [
    ('Nº de Operación', 9, '0', 'Correlativo de la fila', '1'),
    ('Fecha de la Factura', 12, 'DD/MM/YYYY', 'Fecha de emisión', '05/01/2026'),
    ('Nº R.I.F.', 14, '@', 'RIF del proveedor', 'J-12345678-9'),
    ('Proveedor o Razón Social', 34, '@', 'Nombre del proveedor', 'PROVEEDOR EJEMPLO, C.A.'),
    ('N° de Factura', 12, '@', 'Número de la factura', '004520'),
    ('N° Control de Factura', 14, '@', 'Número de control', '00-004520'),
    ('N° Nota de Débito', 12, '@', 'Solo si la fila es una nota de débito', ''),
    ('N° Nota de Crédito', 12, '@', 'Solo si la fila es una nota de crédito', ''),
    ('N° Factura Afectada', 12, '@', 'Factura a la que se aplica la nota', ''),
    ('Total Compras Internas Incluyendo el IVA', 16, '#,##0.00', 'Total del documento', '1.160,00'),
    ('Compras sin Derecho a Crédito IVA', 15, '#,##0.00', 'Monto exento', '0,00'),
    ('Base Imponible', 15, '#,##0.00', 'Base gravada', '1.000,00'),
    ('% Alícuota', 9, '0%', 'Alícuota del IVA (16%)', '16%'),
    ('IVA', 14, '#,##0.00', 'IVA de la factura (crédito fiscal)', '160,00'),
    ('Comprobante de Retención', 17, '@', 'N° del comprobante que le entregaste al proveedor', ''),
    ('IVA Retenido al Vendedor', 15, '#,##0.00', 'IVA que le retuviste', '0,00'),
]
SUMAN = {'Total Ventas Incluyendo el IVA', 'Ventas Internas No Gravadas', 'Base Imponible', 'Impuesto IVA',
         'IGTF', 'IVA Retenido por el Comprador', 'Total Compras Internas Incluyendo el IVA',
         'Compras sin Derecho a Crédito IVA', 'IVA', 'IVA Retenido al Vendedor'}

fino = Side(style='thin', color='C9D4DF')
borde = Border(left=fino, right=fino, top=fino, bottom=fino)


def hoja_mes(ws, titulo, cols, anio, mes):
    n = len(cols)
    ult = get_column_letter(n)
    ws.sheet_view.showGridLines = False
    # 1 · la marca
    ws.merge_cells('A1:%s1' % ult)
    ws['A1'] = 'DigiAccount  ·  ' + titulo
    ws['A1'].font = Font(name=F, size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A1'].alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[1].height = 28
    # 2 · empresa y RIF (los llena el contador)
    ws['A2'] = 'Empresa:'
    ws.merge_cells('B2:F2')
    ws['G2'] = 'RIF:'
    ws.merge_cells('H2:I2')
    for c in ('A2', 'G2'):
        ws[c].font = Font(name=F, size=10, bold=True, color=SLATE)
        ws[c].alignment = Alignment(horizontal='right')
    for c in ('B2', 'H2'):
        ws[c].fill = PatternFill('solid', fgColor=AMARILLO)
        ws[c].font = Font(name=F, size=10, bold=True, color=NAVY)
    # 3 · el período (lo lee el importador)
    ultimo = calendar.monthrange(anio, mes)[1]
    ws.merge_cells('A3:%s3' % ult)
    ws['A3'] = 'Mes %s: 01/%02d/%d al %02d/%02d/%d' % (MESES[mes - 1].capitalize(), mes, anio, ultimo, mes, anio)
    ws['A3'].font = Font(name=F, size=11, bold=True, color=NAVY)
    # 4 · la indicación
    ws.merge_cells('A4:%s4' % ult)
    ws['A4'] = ('Una fila por documento. Anuladas: número y control con los montos en 0. '
                'Para una quincena, cambia el rango de la fila 3. No borres la fila TOTAL.')
    ws['A4'].font = Font(name=F, size=9, italic=True, color=SLATE)
    # 5 · encabezados
    for j, (enc, ancho, _, _, _) in enumerate(cols, start=1):
        c = ws.cell(row=5, column=j, value=enc)
        c.font = Font(name=F, size=9, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(bottom=Side(style='medium', color=CYAN))
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[5].height = 42
    # 6.. renglones
    for i in range(DESDE, DESDE + FILAS):
        for j, (_, _, formato, _, _) in enumerate(cols, start=1):
            c = ws.cell(row=i, column=j)
            c.number_format = formato
            c.font = Font(name=F, size=9)
            c.border = borde
            if (i - DESDE) % 2:
                c.fill = PatternFill('solid', fgColor=CLARO)
    # la fila TOTAL (el importador cuadra contra ella)
    t = DESDE + FILAS
    ws.cell(row=t, column=1, value='TOTAL')
    for j, (enc, _, formato, _, _) in enumerate(cols, start=1):
        c = ws.cell(row=t, column=j)
        c.font = Font(name=F, size=10, bold=True, color=NAVY)
        c.border = Border(top=Side(style='medium', color=NAVY))
        if enc in SUMAN:
            L = get_column_letter(j)
            c.value = '=SUM(%s%d:%s%d)' % (L, DESDE, L, t - 1)
            c.number_format = formato
    # alícuotas válidas
    j_alic = [k for k, col in enumerate(cols, start=1) if col[0] == '% Alícuota'][0]
    dv = DataValidation(type='list', formula1='"0.16,0.08,0.31,0"', allow_blank=True,
                        error='Alícuotas: 16%, 8% (reducida), 31% (general + adicional) o 0.')
    ws.add_data_validation(dv)
    L = get_column_letter(j_alic)
    dv.add('%s%d:%s%d' % (L, DESDE, L, t - 1))
    ws.freeze_panes = 'C6'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '5:5'
    ws.sheet_properties.tabColor = CYAN if mes % 2 else NAVY


def instrucciones(ws, titulo, cols, anio, tipo):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 52
    ws.column_dimensions['D'].width = 24
    ws.merge_cells('B2:D2')
    ws['B2'] = 'DigiAccount  ·  ' + titulo + '  ·  ' + str(anio)
    ws['B2'].font = Font(name=F, size=16, bold=True, color='FFFFFF')
    ws['B2'].fill = PatternFill('solid', fgColor=NAVY)
    ws['B2'].alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[2].height = 34
    pasos = [
        'Cómo usar esta plantilla',
        '1. Llena una hoja por mes (ENERO … DICIEMBRE). Todas tienen el mismo formato.',
        '2. Escribe la empresa y su RIF en la fila 2 de cada hoja que uses.',
        '3. Una fila por documento: factura, nota de crédito o nota de débito.',
        '4. Facturas anuladas: escribe su número y su control con los montos en 0. Así no se pierde el correlativo.',
        '5. Contribuyente especial que declara por quincena: duplica la hoja del mes y cambia el rango de la fila 3 '
        '(por ejemplo «01/01/%d al 15/01/%d»).' % (anio, anio),
        '6. No cambies los títulos de las columnas ni borres la fila TOTAL: con ellos el sistema reconoce el libro '
        'y comprueba que cuadra.',
        '7. En DigiAccount: Módulo Fiscal → Libro de %s → Importar Excel. Antes de guardar verás cada mes con '
        'sus totales para confirmar.' % ('Compras' if tipo == 'compra' else 'Ventas'),
        '8. Las hojas que dejes vacías no se cargan. Los clientes y proveedores se crean solos a partir del RIF.',
    ]
    fila = 4
    for k, p in enumerate(pasos):
        ws.merge_cells('B%d:D%d' % (fila, fila))
        c = ws['B%d' % fila]
        c.value = p
        c.font = Font(name=F, size=12 if k == 0 else 10, bold=(k == 0), color=NAVY if k == 0 else '1A2230')
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[fila].height = 20 if k == 0 else 30
        fila += 1
    fila += 1
    for j, enc in enumerate(['Columna', 'Qué va', 'Ejemplo'], start=2):
        c = ws.cell(row=fila, column=j, value=enc)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
    fila += 1
    for k, (enc, _, _, que, ej) in enumerate(cols):
        for j, v in enumerate([enc, que, ej], start=2):
            c = ws.cell(row=fila, column=j, value=v)
            c.font = Font(name=F, size=10, bold=(j == 2), color=NAVY if j == 2 else '1A2230')
            c.border = borde
            if k % 2:
                c.fill = PatternFill('solid', fgColor=CLARO)
        fila += 1
    fila += 1
    ws.merge_cells('B%d:D%d' % (fila, fila))
    ws['B%d' % fila] = 'DigiAccount · digiaccount.io · Hecho por un Contador Público.'
    ws['B%d' % fila].font = Font(name=F, size=9, italic=True, color=SLATE)
    ws.sheet_properties.tabColor = AMARILLO


def generar(tipo, anio, destino):
    cols = VENTAS if tipo == 'venta' else COMPRAS
    titulo = 'LIBRO DE VENTAS' if tipo == 'venta' else 'LIBRO DE COMPRAS'
    wb = Workbook()
    instrucciones(wb.active, titulo, cols, anio, tipo)
    wb.active.title = 'INSTRUCCIONES'
    for m in range(1, 13):
        hoja_mes(wb.create_sheet('%s %d' % (MESES[m - 1], anio)), titulo, cols, anio, m)
    wb.save(destino)


if __name__ == '__main__':
    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    carpeta = os.path.join(raiz, 'assets', 'plantillas')
    os.makedirs(carpeta, exist_ok=True)
    for anio in [int(a) for a in (sys.argv[1:] or ['2026'])]:
        for tipo, nombre in (('compra', 'compras'), ('venta', 'ventas')):
            ruta = os.path.join(carpeta, 'plantilla-libro-%s-%d.xlsx' % (nombre, anio))
            generar(tipo, anio, ruta)
            print('ok', os.path.relpath(ruta, raiz))
